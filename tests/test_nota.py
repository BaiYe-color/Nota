import copy
import hashlib
import json
import time
from pathlib import Path
import pytest

from evaluation import evaluate,validate_case
from evals.run_baseline import summarize
from fastapi.testclient import TestClient
from contracts import Options,Card
from engine import Context,Cancelled,model_error_message
from storage import Store,Conflict
from server import create_app
from generation import build_concept_graph,card_validator,normalize_knowledge,plan,quality,validate_written,write_chapter,PAPER_OUTLINE
from asr import _normalize_paraformer_result,semantic_segments
from materials import detect_material_type

@pytest.fixture
def client(tmp_path,monkeypatch):
    def fake_llm(self,role,system,data,**kwargs):
        self.check()
        if role=='cards':
            result={'cards':[{'id':str(i),'title':b['text'][:20],'points':[b['text']],
                'source_ids':[b['id']],'importance':3,'formulas':[],'examples':[]} for i,b in enumerate(data)]}
        elif role=='outline' and isinstance(data,dict) and 'cards' in data:
            result={'chapters':[{'id':'tmp','title':'知识总结','card_ids':[b['id'] for b in data['cards']]}]}
        elif role=='outline': result={'issues':[]}
        elif role=='writer':
            result={'markdown':('修订：'+data['revision_request']+'\n\n' if 'revision_request' in data else '')+'\n\n'.join(s['text'] for s in data['sources']),
                    'questions':['关键条件是什么？']}
        else: result={'text':'扫描页识别内容','uncertain':False}
        return kwargs['validate'](result) if kwargs.get('validate') else result
    monkeypatch.setattr(Context,'llm',fake_llm)
    with TestClient(create_app(tmp_path/'data')) as c: yield c

def wait_job(client,jid):
    deadline=time.monotonic()+60
    while time.monotonic()<deadline:
        j=client.get('/api/jobs/'+jid).json()
        if j['status'] not in ('queued','running'): return j
        time.sleep(.02)
    pytest.fail('job did not finish')

def upload(client,name='lesson.txt',content=b'Important definition: an index maps terms to documents.'):
    r=client.post('/api/upload',files={'files':(name,content)})
    assert r.status_code==200,r.text
    return r.json()['files'][0]['id']

def create_note(client):
    fid=upload(client)
    r=client.post('/api/generate',json={'file_ids':[fid],'options':{}})
    assert r.status_code==200,r.text
    j=wait_job(client,r.json()['job_id'])
    assert j['status']=='succeeded',j
    return client.get('/api/notes/'+j['result']['note_id']).json(),j

def test_end_to_end_edit_revise_restore_export(client):
    note,job=create_note(client); nid=note['id']; cid=note['data']['chapters'][0]['id']
    assert note['data']['quality']['covered_cards']==1
    original=note['data']['chapters'][0]['markdown']
    assert client.put(f'/api/notes/{nid}/chapters/{cid}',json={'expected_version':1,'markdown':'User edited material'}).json()['version']==2
    assert client.put(f'/api/notes/{nid}/chapters/{cid}',json={'expected_version':1,'markdown':'stale'}).status_code==409
    r=client.post(f'/api/notes/{nid}/chapters/{cid}/revise',json={'expected_version':2,'instruction':'更精简'})
    j=wait_job(client,r.json()['job_id']);assert j['status']=='succeeded',j
    current=client.get('/api/notes/'+nid).json(); assert current['version']==3
    assert '更精简' in current['data']['chapters'][0]['markdown']
    assert 'User edited material' in client.get(f'/api/notes/{nid}?version=2').json()['data']['chapters'][0]['markdown']
    diff=client.get(f'/api/notes/{nid}/diff?before=1&after=2').json()['diff'];assert 'User edited material' in diff
    assert client.post(f'/api/notes/{nid}/restore',json={'expected_version':3,'version':1}).json()['version']==4
    assert client.get('/api/notes/'+nid).json()['data']['chapters'][0]['markdown']==original
    for fmt in ('md','html','zip','pdf'):
        r=client.post(f'/api/notes/{nid}/export?format={fmt}')
        j=wait_job(client,r.json()['job_id']);assert j['status']=='succeeded',j
        response=client.get('/api/artifacts/'+j['result']['artifact']);assert response.status_code==200
        assert len(response.content)>30

def test_same_filename_different_content_isolated(client):
    a=upload(client,content=b'Alpha document with original contents')
    b=upload(client,content=b'Beta document with different contents')
    assert a!=b
    assert client.get('/api/files/'+a).content.startswith(b'Alpha')
    assert client.get('/api/files/'+b).content.startswith(b'Beta')
    assert upload(client,content=b'Alpha document with original contents')==a

def test_missing_upload_is_reported_and_reupload_repairs_copy(client):
    content=b'Important definition: an index maps terms to documents.'
    fid=upload(client,'lesson.txt',content)
    path=Path(client.app.state.store.file(fid)['path'])
    path.unlink()
    missing=client.post('/api/generate',json={'file_ids':[fid],'options':{}})
    assert missing.status_code==400
    assert '原始文件已不存在，请重新上传' in missing.json()['detail']
    assert client.get('/api/files/'+fid).status_code==410
    assert upload(client,'lesson.txt',content)==fid
    assert path.read_bytes()==content

@pytest.mark.parametrize('name',['../evil.txt','..\\evil.txt','C:evil.txt','evil.html','evil.exe'])
def test_unsafe_upload_rejected(client,name):
    assert client.post('/api/upload',files={'files':(name,b'data')}).status_code==400

def test_paths_and_cross_site_rejected(client):
    assert client.post('/api/generate',json={'paths':['C:/secret.txt']}).status_code==422
    assert client.post('/api/generate',json={'file_ids':['missing']}).status_code==404
    assert client.post('/api/preferences',headers={'Origin':'https://evil.example'},json={}).status_code==403
    assert client.get('/api/health',headers={'Host':'evil.example'}).status_code==400
    assert client.get('/api/file?path=C:/Windows/win.ini').status_code==404

def test_empty_duplicate_and_limit(client,monkeypatch):
    assert client.post('/api/upload',files={'files':('empty.txt',b'')}).status_code==400
    fid=upload(client)
    assert client.post('/api/generate',json={'file_ids':[fid,fid]}).status_code==400
    monkeypatch.setenv('NOTA_MAX_UPLOAD_MB','0')
    assert client.post('/api/upload',files={'files':('x.txt',b'data')}).status_code==413

def test_sse_replay_cursor(client):
    _,job=create_note(client)
    text=client.get('/api/stream/'+job['id']).text
    assert 'event: end' in text and 'id:' in text
    ids=[int(line[4:]) for line in text.splitlines() if line.startswith('id: ')]
    text2=client.get('/api/stream/'+job['id'],headers={'Last-Event-ID':str(ids[-1])}).text
    assert '"kind": "progress"' not in text2
    assert 'event: end' in text2

def test_cache_version_options_and_failure(tmp_path):
    store=Store(tmp_path);jid=store.new_job('test',{});ctx=Context(store,jid);calls=[]
    def run(): calls.append(1);return {'ok':True}
    assert ctx.cached('sample',{'sha':'a','opt':False},run)=={'ok':True}
    ctx.cached('sample',{'sha':'a','opt':False},run)
    ctx.cached('sample',{'sha':'a','opt':True},run)
    ctx.cached('sample',{'sha':'b','opt':True},run)
    assert len(calls)==3
    def fail(): raise ValueError('failed')
    with pytest.raises(ValueError):ctx.cached('sample',{'sha':'fail'},fail)
    ctx.cached('sample',{'sha':'fail'},run);assert len(calls)==4
    store.cancel(jid)
    with pytest.raises(Cancelled):ctx.cached('sample',{'sha':'a','opt':False},run)

def test_restart_recovery(tmp_path):
    a=Store(tmp_path);jid=a.new_job('generate',{});a.set_job(jid,'running')
    b=Store(tmp_path);b.recover();assert b.job(jid)['status']=='interrupted'

def test_recent_notes_keep_latest_ten_and_return_latest_versions(client):
    store = client.app.state.store
    ids = []
    for index in range(11):
        nid = store.create_note(f'job-{index}', {'title': f'笔记 {index}'})
        ids.append(nid)
        time.sleep(.002)
    store.save_revision(ids[0], 1, {'title': '笔记 0（更新版）'}, '编辑')

    response = client.get('/api/notes?limit=99')
    assert response.status_code == 200
    notes = response.json()
    assert len(notes) == 10
    assert notes[0]['id'] == ids[0]
    assert notes[0]['title'] == '笔记 0（更新版）'
    assert notes[0]['version'] == 2
    assert ids[1] not in {note['id'] for note in notes}


def test_delete_note_removes_its_history_but_keeps_other_notes(client):
    store = client.app.state.store
    first = store.create_note('job-first', {'title': '保留的笔记'})
    removed = store.create_note('job-removed', {'title': '将删除的笔记'})
    store.save_revision(removed, 1, {'title': '将删除的笔记 v2'}, '编辑')

    response = client.delete(f'/api/notes/{removed}')
    assert response.status_code == 200
    assert client.get(f'/api/notes/{removed}').status_code == 404
    assert [note['id'] for note in client.get('/api/notes').json()] == [first]
    assert client.delete(f'/api/notes/{removed}').status_code == 404


def test_atomic_revision_conflict(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    store=Store(tmp_path);nid=store.create_note('job',{'value':0})
    def save(value):
        try:return store.save_revision(nid,1,{'value':value},'edit')
        except Conflict:return 'conflict'
    with ThreadPoolExecutor(2) as pool:results=list(pool.map(save,[1,2]))
    assert sorted(map(str,results))==['2','conflict']
    assert len(store.history(nid))==2

def test_unknown_sources_and_bad_markdown():
    validate=card_validator([{'id':'known'}])
    with pytest.raises(ValueError):validate({'cards':[{'id':'a','title':'x','points':['y'],'source_ids':['invented']}]})
    with pytest.raises(ValueError):validate_written({'markdown':'```broken'})
    assert '![' not in validate_written({'markdown':'Hi ![x](file:///secret)'})['markdown']
    with pytest.raises(ValueError,match='公式占位符'):
        validate_written({'markdown':'$$formula-cartesian$$'})
    formula_input={'cards':[{'id':'a','title':'x','points':['y'],
        'source_ids':['known'],'importance':3,'formulas':['formula-1'],'examples':[]}]}
    formula_validator=card_validator([{'id':'known'}])
    with pytest.raises(ValueError,match='重新生成完整 JSON'):
        formula_validator(formula_input)
    result=formula_validator(formula_input)
    assert result['knowledge']['formulas']==[]
    from rendering import guard_math
    guard_math(r'$$R \Join S$$')
    assert validate_written({'markdown':'正文','questions':[{'question':'关键条件是什么？','answer':'x'}]})['questions']==['关键条件是什么？']

def test_formula_prompt_treats_uncertain_math_as_optional():
    from generation import CARDS
    assert '公式是可选证据' in CARDS and 'cards.formulas 始终返回 []' in CARDS

def test_preview_sanitization():
    from rendering import safe_html
    result=safe_html('<script>alert(1)</script>\n\n<img src="x" onerror="x">\n\n[x](javascript:alert(1))')
    assert '<script' not in result and '<img' not in result and 'href="javascript:' not in result

def test_asr_normalization():
    result=_normalize_paraformer_result({'transcripts':[{'sentences':[{'begin_time':2000,'end_time':3000,'text':'B','speaker_id':2},{'begin_time':0,'end_time':1000,'text':'A','speaker_id':1},{'text':' '}]}]})
    assert [s['text'] for s in result]==['A','B']
    assert result[0]['speaker_id']==1
    with pytest.raises(ValueError):_normalize_paraformer_result({'transcripts':[{'sentences':[{'begin_time':20,'end_time':10,'text':'bad'}]}]})

def test_asr_semantic_segments_merge_turns_and_split_topic_boundaries():
    segments=[{'start_ms':0,'end_ms':800,'text':'We first review inverted indexes.','speaker_id':1},
              {'start_ms':900,'end_ms':1600,'text':'They map terms to documents.','speaker_id':1},
              {'start_ms':1800,'end_ms':2400,'text':'Next we discuss Boolean retrieval.','speaker_id':1},
              {'start_ms':2500,'end_ms':3000,'text':'I have a question.','speaker_id':2}]
    grouped=semantic_segments(segments)
    assert len(grouped)==3 and 'map terms' in grouped[0]['text']
    assert grouped[1]['text'].startswith('Next') and grouped[2]['speaker_id']==2

def test_transcript_correction_creates_revision_and_regenerates(client,monkeypatch):
    store=client.app.state.store
    jid=store.new_job('generate',{'file_ids':[],'options':{}})
    data={'title':'recording','options':Options().model_dump(),'sources':[{
        'id':'audio-1','document_id':'file-1','document_name':'recording.wav',
        'start_ms':0,'end_ms':2200,'text':'old transcript','kind':'audio','method':'asr',
        'speaker_id':1,'asr_sentences':[]}], 'counts':[], 'cards':[], 'knowledge':{},
        'chapters':[], 'materials':[], 'transcript':'[00:00] old transcript',
        'warnings':[], 'metrics':[], 'pipeline_version':'test','models':{},
        'quality':{'issues':[]}}
    nid=store.create_note(jid,data)
    saved=client.put(f'/api/notes/{nid}/transcript',json={
        'expected_version':1,'segments':[{'id':'audio-1','text':'corrected transcript'}]})
    assert saved.status_code==200 and saved.json()['version']==2
    note=client.get(f'/api/notes/{nid}').json()['data']
    assert note['sources'][0]['text']=='corrected transcript'
    assert note['transcript_corrected'] is True
    captured={}
    monkeypatch.setattr(client.app.state.service,'submit',lambda kind,payload:captured.update(kind=kind,payload=payload) or 'job-2')
    rerun=client.post(f'/api/notes/{nid}/transcript/regenerate?expected_version=2')
    assert rerun.json()['job_id']=='job-2'
    assert captured=={'kind':'regenerate','payload':{'note_id':nid,'expected_version':2}}

def test_retry_can_skip_review_questions(client,monkeypatch):
    store=client.app.state.store;service=client.app.state.service
    jid=store.new_job('generate',{'file_ids':['test-file'],'options':{'include_review_questions':True}})
    store.set_job(jid,'failed',error='writer failed')
    captured={}
    monkeypatch.setattr(service,'submit',lambda kind,payload:captured.update(kind=kind,payload=payload) or 'replacement')
    result=client.post('/api/jobs/'+jid+'/retry',json={'skip_review_questions':True})
    assert result.status_code==200 and result.json()['job_id']=='replacement'
    assert captured['kind']=='generate' and captured['payload']['options']['include_review_questions'] is False
    assert captured['payload']['_resume_from']==jid
    result=client.post('/api/jobs/'+jid+'/retry',json={'degraded_mode':True})
    assert result.status_code==200 and captured['payload']['options']['degraded_mode'] is True
    assert captured['payload']['options']['include_images'] is False
    assert '_resume_from' not in captured['payload']

def test_transcript_only_no_writer(client,monkeypatch):
    import extraction
    monkeypatch.setattr(extraction,'transcribe_local_file',lambda *a,**k:[{'start_ms':0,'end_ms':1000,'text':'今天决定先测试，不发布。','speaker_id':1}])
    fid=upload(client,'test.wav',b'fake test audio')
    j=wait_job(client,client.post('/api/generate',json={'file_ids':[fid],'options':{'output_form':'transcript','diarization':True}}).json()['job_id'])
    assert j['status']=='succeeded'
    d=client.get('/api/notes/'+j['result']['note_id']).json()['data'];assert not d['chapters'];assert '说话人1' in d['transcript'];assert '不发布' in d['transcript']

def test_preferences_persist(client):
    r=client.put('/api/preferences',json={'style':'短句与箭头','template':'revision'})
    assert r.status_code==200
    assert client.get('/api/preferences').json()['style']=='短句与箭头'

def test_model_settings_never_return_keys_or_store_them_in_job_payload(client):
    payload={'default':{'model':'test-model','base_url':'https://example.invalid/v1','api_key':'private-key'},
             'cards':{'use_default':True},'outline':{'use_default':True},'writer':{'use_default':True},'vision':{'use_default':True},
             'asr':{'model':'paraformer-v2','base_url':'https://example.invalid/api/v1','api_key':'asr-private'}}
    saved=client.put('/api/model-settings',json=payload)
    assert saved.status_code==200 and saved.json()['default']['configured'] is True
    public=client.get('/api/model-settings').json()
    assert 'private-key' not in str(public) and public['asr']['configured'] is True
    jid=client.app.state.service.submit('generate',{'file_ids':[],'materials':[],'options':{}})
    job=client.app.state.store.job(jid)
    assert job['payload']['model_config_id'] and 'private-key' not in str(job['payload'])

def test_pdf_native_and_failed_page_checkpoint(tmp_path,monkeypatch):
    import pymupdf
    from extraction import extract
    import engine
    doc=pymupdf.open()
    for text in ['First page with a long native text sentence about sorted arrays.', 'Second page also contains enough native text to avoid OCR.']:
        p=doc.new_page();p.insert_text((40,60),text)
    path=tmp_path/'native.pdf';doc.save(path);doc.close()
    store=Store(tmp_path/'data');jid=store.new_job('generate',{});ctx=Context(store,jid)
    f={'id':'a'*64,'sha':'hash','name':'native.pdf','path':str(path)}
    monkeypatch.setattr(ctx,'llm',lambda *a,**k:pytest.fail('Native PDF should not call vision'))
    blocks,counts=extract(ctx,[f],Options(include_images=False))
    assert counts[0]['processed']==2 and all(b['method']=='native' for b in blocks)
    assert {b['page'] for b in blocks}=={1,2}
    ctx2=Context(store,jid);extract(ctx2,[f],Options(include_images=False))
    assert len([m for m in ctx2.metrics if m.get('cache_hit')])==2

def test_ppt_native_notes_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    from extraction import extract
    prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[1]);s.shapes.title.text='Definition';s.placeholders[1].text='Input must be sorted.';s.notes_slide.notes_text_frame.text='Keep the boundary condition.'
    table=s.shapes.add_table(1,2,Inches(1),Inches(4),Inches(4),Inches(1)).table;table.cell(0,0).text='Key';table.cell(0,1).text='Value'
    p=tmp_path/'native.pptx';prs.save(p)
    store=Store(tmp_path/'data');ctx=Context(store,store.new_job('generate',{}))
    blocks,_=extract(ctx,[{'id':'b'*64,'sha':'ppt','name':'native.pptx','path':str(p)}],Options())
    text=''.join(b['text'] for b in blocks)
    assert 'boundary condition' in text and 'Key | Value' in text

def test_asr_upload_protocol_and_headers(tmp_path,monkeypatch):
    import asr
    monkeypatch.setenv('DASHSCOPE_API_KEY','test-only')
    calls=[]
    class Response:
        ok=True;status_code=200
        def __init__(self,data):self.data=data
        def json(self):return self.data
    def get(url,**kwargs):
        calls.append(('GET',url,kwargs))
        if url.endswith('/uploads'):return Response({'data':{'upload_dir':'test-dir','upload_host':'https://test.example/upload','oss_access_key_id':'test','signature':'test','policy':'test','x_oss_object_acl':'private','x_oss_forbid_overwrite':'true'}})
        if '/tasks/' in url:return Response({'output':{'task_status':'SUCCEEDED','results':[{'transcription_url':'https://test.example/result'}]}})
        return Response({'transcripts':[{'sentences':[{'begin_time':0,'end_time':1000,'text':'测试','speaker_id':0}]}]})
    def post(url,**kwargs):
        calls.append(('POST',url,kwargs))
        return Response({'output':{'task_id':'test-task'}})
    monkeypatch.setattr(asr.requests,'get',get);monkeypatch.setattr(asr.requests,'post',post)
    p=tmp_path/'audio.wav';p.write_bytes(b'test')
    segs=asr.transcribe_local_file(p,enable_diarization=True)
    request=next(c for c in calls if c[0]=='POST' and c[1].endswith('/transcription'))[2]
    assert request['json']['input']['file_urls'][0].startswith('oss://')
    assert request['headers']['X-DashScope-OssResourceResolve']=='enable'
    assert request['json']['parameters']['diarization_enabled'] is True
    assert segs[0]['speaker_id']==0

def test_cancelled_job_no_pipeline_calls(client,monkeypatch):
    import service
    store=client.app.state.store;svc=client.app.state.service
    jid=store.new_job('generate',{'file_ids':[],'options':{}});store.cancel(jid)
    monkeypatch.setattr(service,'extract',lambda *a,**k:pytest.fail('cancelled job must not parse'))
    svc.run(jid);assert store.job(jid)['status']=='cancelled'

def test_model_fallback_records_actual_model(tmp_path,monkeypatch):
    import engine
    engine._UNAVAILABLE_UNTIL.clear()
    monkeypatch.setenv('DEEPSEEK_API_KEY','test-only');monkeypatch.setenv('NOTA_TEXT_FALLBACK_MODEL','deepseek-chat')
    monkeypatch.setitem(engine.MODELS,'cards','gpt-unavailable')
    store=Store(tmp_path);ctx=Context(store,store.new_job('test',{}));seen=[]
    def call(role,system,data,model,**kwargs):
        seen.append(model)
        if model=='gpt-unavailable':raise engine.ModelUnavailable('down')
        return {'model':model}
    monkeypatch.setattr(ctx,'_llm_for_model',call)
    assert ctx.llm('cards','prompt',{})=={'model':'deepseek-chat'}
    assert seen==['gpt-unavailable','deepseek-chat']
    assert any('deepseek-chat' in e.get('message','') for e in store.events(ctx.jid))
    engine._UNAVAILABLE_UNTIL.clear()

def test_math_export_guard():
    from rendering import guard_math
    guard_math(r'$x=\frac{1}{2}$')
    for attack in [r'$\input{secret}$',r'$\csname input\endcsname{secret}$',r'$^^5cinput{secret}$',r'$\begin{filecontents}bad\end{filecontents}$']:
        with pytest.raises(ValueError):guard_math(attack)

def test_only_target_chapter_revision_changed(client):
    note,_=create_note(client);store=client.app.state.store
    data=copy.deepcopy(note['data']);second=copy.deepcopy(data['chapters'][0]);second['id']='second';second['title']='Keep';second['markdown']='DO NOT CHANGE';data['chapters'].append(second)
    store.save_revision(note['id'],1,data,'fixture')
    cid=data['chapters'][0]['id'];r=client.post(f"/api/notes/{note['id']}/chapters/{cid}/revise",json={'expected_version':2,'instruction':'压缩'})
    j=wait_job(client,r.json()['job_id']);assert j['status']=='succeeded'
    current=client.get('/api/notes/'+note['id']).json()
    assert current['data']['chapters'][1]==second|{'html':'<p>DO NOT CHANGE</p>'}

def test_model_data_validation_never_caches_failed_result(tmp_path,monkeypatch):
    from engine import Context
    # Cache callback fails before atomic commit; later valid run must execute again.
    store=Store(tmp_path);ctx=Context(store,store.new_job('test',{}))
    with pytest.raises(ValueError):ctx.cached('validated',{},lambda:(_ for _ in ()).throw(ValueError('bad output')))
    assert not list((tmp_path/'cache').rglob('*.json'))

def test_second_server_cannot_interrupt_live_jobs(tmp_path):
    store=Store(tmp_path/'locked');jid=store.new_job('test',{});store.set_job(jid,'running')
    with store.runtime_lock():
        with pytest.raises(RuntimeError):
            with TestClient(create_app(tmp_path/'locked')):pass
        assert store.job(jid)['status']=='running'
    with store.runtime_lock():pass

def test_math_preview_preserves_latex():
    from rendering import safe_html
    text=r'$$\begin{matrix}a_1 & b_2 \\ c_3 & d_4\end{matrix}$$'
    result=safe_html(text)
    assert r'\begin{matrix}' in result and r'\\ c_3' in result and '<em>' not in result

def test_pdf_figure_is_cropped_visual_evidence(tmp_path,monkeypatch):
    import io
    import pymupdf
    from PIL import Image,ImageDraw
    from extraction import extract,pdf_visual_regions
    doc=pymupdf.open();page=doc.new_page()
    page.insert_textbox(pymupdf.Rect(50,40,560,120),'Figure 1 explains the encoder and decoder architecture. '+('Native explanation. '*8),fontsize=11)
    picture=Image.new('RGB',(500,260),'white');draw=ImageDraw.Draw(picture)
    draw.rectangle((20,20,480,240),outline='black',width=4);draw.text((60,110),'ENCODER -> ATTENTION -> DECODER',fill='black')
    image=io.BytesIO();picture.save(image,format='PNG')
    figure_rect=pymupdf.Rect(120,180,490,400);page.insert_image(figure_rect,stream=image.getvalue())
    page.insert_text((150,425),'Figure 1: Model architecture')
    path=tmp_path/'figure.pdf';doc.save(path);doc.close()
    with pymupdf.open(path) as check:
        regions=pdf_visual_regions(check[0],{check[0].get_images()[0][0]:1},1)
        assert regions and regions[0][1].get_area()<check[0].rect.get_area()*.5
    store=Store(tmp_path/'data');ctx=Context(store,store.new_job('generate',{}))
    def fake(role,system,data,**kwargs):
        result={'text':'Figure 1：编码器、注意力与解码器结构\n- 展示信息流向','uncertain':False}
        return kwargs['validate'](result)
    monkeypatch.setattr(ctx,'llm',fake)
    blocks,_=extract(ctx,[{'id':'c'*64,'sha':'figure-hash','name':'figure.pdf','path':str(path)}],Options(include_images=True))
    native=[b for b in blocks if b['method']=='native'];visual=[b for b in blocks if b['method']=='vision-region']
    assert native and all(not b['asset'] for b in native)
    assert len(visual)==1 and visual[0]['asset'] and visual[0]['kind']=='visual'
    with pymupdf.open(store.root/'artifacts'/visual[0]['asset']) as cropped:
        assert cropped[0].rect.width<1000

def test_exported_figure_has_semantic_caption():
    from rendering import note_markdown
    source={'id':'figure','document_name':'paper.pdf','page':3,'text':'Figure 1：Transformer 编码器与解码器结构','asset':'figure.png'}
    data={'title':'论文笔记','sources':[source],'options':{'include_images':True},
          'chapters':[{'title':'Transformer 架构','markdown':'编码器与解码器通过注意力连接。','status':'ready','issues':[],
                       'source_ids':['figure']} ]}
    rendered=note_markdown(data)
    assert '### 与本节对应的图表' not in rendered
    assert '![图 1：' in rendered
    assert 'Transformer 编码器与解码器结构' in rendered
    assert 'images/figure.png' in rendered

def test_ir_formula_and_visual_tokens_are_resolved_and_internal_review_removed():
    from generation import clean_final_markdown
    markdown='结论（cl-a1b2c3）。\n\n[[FORMULA:fm-1]]\n\n[[VISUAL:vi-1]]\n\n待核对项：内部检查。\n\n### 核对提示\n\n- 建议改为旧文本'
    result=clean_final_markdown(markdown,[{'id':'fm-1','latex':r'\beta_1=0.9'}],[{'id':'vi-1'}])
    assert '$$\\beta_1=0.9$$' in result and '[[VISUAL:vi-1]]' in result
    assert '核对提示' not in result and '建议改为' not in result and '待核对项' not in result and 'cl-a1b2c3' not in result
    assert clean_final_markdown('[[FORMULA:bad]]',[{'id':'bad','latex':'formula-1'}],[])==''

def test_note_markdown_hides_source_labels_unless_requested():
    from rendering import note_markdown
    data={'title':'Test','sources':[{'id':'s1','document_name':'source.pdf','page':1,'text':'source'}],
          'options':{'include_images':False},'knowledge':{},
          'chapters':[{'title':'One','status':'ready','markdown':'正文','source_ids':['s1'],'visual_ids':[]}]}
    assert '来源：' not in note_markdown(data)
    assert '来源：source.pdf · 第 1 页/单元' in note_markdown(data,include_sources=True)

def test_final_markdown_compacts_excessive_subheadings():
    from generation import clean_final_markdown
    markdown='\n\n'.join('### 主题 '+str(n)+'\n\n内容 '+str(n) for n in range(1,6))
    result=clean_final_markdown(markdown,[],[])
    assert result.count('### ')==3 and '**主题 4**' in result and '**主题 5**' in result

def test_global_finalizer_removes_repeated_paragraph_with_same_numbers():
    from generation import finalize_chapters
    paragraph='自注意力把任意两个位置的最大路径长度降到 O(1)，从而改善长距离依赖并支持并行计算。'
    chapters=[{'status':'ready','markdown':paragraph},{'status':'ready','markdown':paragraph+'\n\n另一条独立结论。'}]
    result,removed=finalize_chapters(chapters,{})
    assert removed==1 and paragraph not in result[1]['markdown'] and '另一条独立结论' in result[1]['markdown']

def test_visual_caption_removes_source_figure_number_and_incomplete_tail():
    from rendering import visual_caption
    source={'document_name':'paper.pdf','page':1,'text':''}
    visual={'title':'Figure 5 illustrates the attention mechanisms within an encoder layer, by visualizing how different heads process sen'}
    caption=visual_caption(source,visual)
    assert not caption.startswith('Figure 5') and not caption.endswith('process sen') and len(caption)<=100

def test_pdf_artifact_warns_but_keeps_image_only_page(tmp_path):
    import io
    import pymupdf
    from PIL import Image
    from rendering import validate_pdf_artifact
    image=io.BytesIO();Image.new('RGB',(200,200),'white').save(image,format='PNG')
    path=tmp_path/'bad.pdf';doc=pymupdf.open();page=doc.new_page();page.insert_image(page.rect,stream=image.getvalue());doc.save(path);doc.close()
    warnings=validate_pdf_artifact(path)
    assert warnings and '图片或公式为主' in warnings[0]
    assert path.exists()

def test_writer_only_receives_current_evidence(tmp_path,monkeypatch):
    store=Store(tmp_path/'data');ctx=Context(store,store.new_job('generate',{}));seen=[]
    cards=[{'id':'current','title':'Current topic','points':['current fact'],'source_ids':['s1'],'importance':3,'formulas':[],'examples':[]},
           {'id':'other','title':'Other topic','points':['SECRET OTHER FACT'],'source_ids':['s2'],'importance':3,'formulas':[],'examples':[]}]
    blocks=[{'id':'s1','text':'current evidence','material_type':'paper','material_role':'primary'},
            {'id':'s2','text':'other evidence','material_type':'general','material_role':'primary'}]
    def fake(role,system,data,**kwargs):
        if role=='writer':
            seen.append(data)
            result={'markdown':'current evidence','questions':['why?']}
        else: result={'issues':[]}
        return kwargs['validate'](result)
    monkeypatch.setattr(ctx,'llm',fake)
    result=write_chapter(ctx,{'id':'ch','title':'Current','card_ids':['current']},cards,blocks,Options())
    assert result['status']=='ready'
    payload=seen[0]
    assert 'outline' not in payload and len(payload['sources'])==1
    assert payload['sources'][0]['id']=='s1' and payload['sources'][0]['text']=='current evidence'
    assert len(payload['material_guidance'])==1 and '研究问题' in payload['material_guidance'][0]
    assert payload['paper_mode'] is False
    assert payload['other_chapter_topics_do_not_repeat']==['Other topic']
    assert 'SECRET OTHER FACT' not in str(payload)


def test_paper_outline_uses_research_evidence_contract():
    seen=[]
    cards=[{'id':'c1','title':'研究问题','points':['x'],'source_ids':['s1'],'importance':3,'claim_ids':[]},
           {'id':'c2','title':'实验结果','points':['y'],'source_ids':['s2'],'importance':3,'claim_ids':[]}]
    class Ctx:
        def llm(self,role,system,data,validate):
            seen.append(system)
            return validate({'chapters':[{'id':'tmp','title':'问题与实验','card_ids':['c1','c2']}]})
        def warn(self,message): raise AssertionError(message)
    chapters=plan(Ctx(),cards,material_type='paper')
    assert chapters[0]['card_ids']==['c1','c2'] and seen==[PAPER_OUTLINE]


def test_paper_primary_keeps_paper_mode_with_supplementary_context(tmp_path,monkeypatch):
    store=Store(tmp_path/'data');ctx=Context(store,store.new_job('generate',{}));seen=[]
    cards=[{'id':'paper','title':'Paper','points':['fact'],'source_ids':['p'],'importance':3,'formulas':[],'examples':[]},
           {'id':'supp','title':'Supplement','points':['context'],'source_ids':['s'],'importance':3,'formulas':[],'examples':[]}]
    blocks=[{'id':'p','text':'paper evidence','material_type':'paper','material_role':'primary'},
            {'id':'s','text':'supplementary evidence','material_type':'textbook','material_role':'supplementary'}]
    def fake(role,system,data,**kwargs):
        if role=='writer':
            seen.append(data); result={'markdown':'paper evidence','questions':[]}
        else: result={'issues':[]}
        return kwargs['validate'](result)
    monkeypatch.setattr(ctx,'llm',fake)
    write_chapter(ctx,{'id':'ch','title':'Paper','card_ids':['paper','supp']},cards,blocks,Options())
    assert seen[0]['paper_mode'] is True

def test_offline_evaluation_reports_claim_pair_and_visual_results():
    note={'id':'note-1','version':2,'data':{
        'chapters':[{'id':'ch-1','title':'注意力架构','markdown':'模型使用 8 个头。',
                     'source_ids':['visual-1']}],
        'sources':[{'id':'visual-1','text':'Figure 2 Multi-Head Attention','asset':'figure.png'}],
        'quality':{'issues':[{'message':'sample'}]},
        'metrics_summary':{'seconds':2.5,'model_calls':3,'input_tokens':100,'output_tokens':20}}}
    case={'id':'sample','source_name':'sample.pdf','source_path':'sample.pdf',
          'material_type':'paper','synthetic':True,
          'required_claims':[{'id':'heads','all':['8 个头'],'any':['多头|注意力']}],
          'forbidden_pairs':[{'id':'wrong-pair','left':'41\\.8','right':'英译德','window':80}],
          'visual_expectations':[{'id':'attention-figure','source_any':['Figure 2'],
                                  'chapter_any':['注意力']} ]}
    report=evaluate(note,case)
    assert report['scores']=={'required_claim_recall':1.0,'forbidden_pair_precision':1.0,
                              'visual_binding_recall':1.0,'unresolved_issue_count':1}
    assert report['cost']['model_calls']==3
    with pytest.raises(ValueError):
        validate_case({'id':'broken','source_name':'x','required_claims':[{'id':'bad','all':['[']} ]})

def test_offline_evaluation_allows_explicit_counterexample_context():
    note={'data':{'chapters':[{'title':'易错点','markdown':'常见错误：将测试集用于参数训练（违反数据隔离要求）。'}],
                  'sources':[]}}
    case={'id':'negative-context','source_name':'x.docx','source_path':'x.docx','material_type':'docx','synthetic':True,
          'forbidden_pairs':[{'id':'leak','left':'测试集','right':'用于参数训练','window':40,
                              'allow_if':['常见错误.{0,20}测试集用于参数训练']}],
          'required_claims':[],'visual_expectations':[],'allowed_omissions':[]}
    assert evaluate(note,case)['scores']['forbidden_pair_precision']==1.0

def test_quality_suite_has_two_reviewed_cases_per_material_type():
    project=Path(__file__).resolve().parents[1]
    cases=[validate_case(json.loads(path.read_text(encoding='utf-8')))
           for path in sorted((project/'evals'/'cases').glob('*.json'))]
    counts={kind:sum(case['material_type']==kind for case in cases)
            for kind in ('paper','scanned_pdf','lecture_slides','docx','meeting_audio')}
    assert counts=={kind:2 for kind in counts}
    assert len({case['id'] for case in cases})==10
    assert all((project/case['source_path']).is_file() for case in cases)

def test_baseline_summary_groups_scores_and_cost_by_material():
    reports=[{'material_type':'paper','scores':{'required_claim_recall':1.0,
              'forbidden_pair_precision':1.0,'visual_binding_recall':.5,'unresolved_issue_count':2},
              'cost':{'seconds':3,'model_calls':2,'input_tokens':100,'output_tokens':20}},
             {'material_type':'docx','scores':{'required_claim_recall':.5,
              'forbidden_pair_precision':1.0,'visual_binding_recall':1.0,'unresolved_issue_count':0},
              'cost':{'seconds':1,'model_calls':1,'input_tokens':30,'output_tokens':10}}]
    result=summarize(reports)
    assert result['average_scores']['required_claim_recall']==.75
    assert result['cost']=={'seconds':4,'model_calls':3,'input_tokens':130,'output_tokens':30}
    assert result['by_material_type']['paper']['unresolved_issue_count']==2

def test_material_type_role_and_style_reference_are_preserved(client):
    primary=upload(client,'paper.txt',b'Primary fact: the measured value is 42.')
    style=upload(client,'my-notes.md','短句。→ 箭头连接。重点加粗。'.encode('utf-8'))
    payload={'file_ids':[primary,style],
             'materials':[{'file_id':primary,'source_type':'paper','role':'primary'},
                          {'file_id':style,'source_type':'personal_notes','role':'style_reference'}],
             'options':{}}
    response=client.post('/api/generate',json=payload)
    assert response.status_code==200,response.text
    job=wait_job(client,response.json()['job_id']);assert job['status']=='succeeded',job
    note=client.get('/api/notes/'+job['result']['note_id']).json()['data']
    assert [item['role'] for item in note['materials']]==['primary','style_reference']
    assert note['materials'][0]['effective_type']=='paper'
    assert {source['document_name'] for source in note['sources']}=={'paper.txt'}
    assert all(source['material_type']=='paper' and source['material_role']=='primary' for source in note['sources'])
    assert '箭头连接' in note['options']['style']
    assert any(row['document']=='my-notes.md' and row['unit']=='风格参考' for row in note['counts'])

def test_material_auto_detection_and_user_override(tmp_path):
    path=tmp_path/'lecture.pptx';path.write_bytes(b'placeholder')
    detected,confidence,reason=detect_material_type({'name':path.name,'path':str(path)},Options())
    assert detected=='lecture_slides' and confidence>.9 and reason

def test_knowledge_ir_has_stable_typed_claims_and_traceability():
    sources=[{'id':'s1','text':'Accuracy is 92% on TestSet.','kind':'table','asset':'table.png'}]
    validate=card_validator(sources)
    raw={'cards':[{'id':'card','title':'Evaluation','points':['Accuracy is 92%.'],'source_ids':['s1'],
                   'importance':4,'formulas':['a=b'],'examples':[],'claim_ids':['claim']}],
         'claims':[{'id':'claim','statement':'Accuracy is 92% on TestSet.','claim_type':'result',
                    'source_ids':['s1'],'conditions':['TestSet'],'confidence':1,'importance':4}],
         'concepts':[{'id':'concept','name':'Accuracy','definition':'Evaluation metric',
                      'claim_ids':['claim'],'depends_on':[]}],
         'formulas':[{'id':'formula','latex':'a=b','meaning':'sample','variables':['a: prediction'],
                      'source_ids':['s1'],'claim_ids':['claim']}],
         'experiments':[{'id':'experiment','name':'Main result','dataset':'TestSet','setup':['fixed split'],
                         'metrics':['Accuracy'],'results':['92%'],'source_ids':['s1'],'claim_ids':['claim']}],
         'visuals':[{'id':'visual','source_id':'s1','title':'Results table','visual_type':'table',
                     'description':'Reports accuracy.','claim_ids':['claim']}]}
    result=validate(raw);knowledge=result['knowledge'];claim=knowledge['claims'][0]
    assert claim['id'].startswith('cl-') and claim['claim_type']=='result' and claim['source_ids']==['s1']
    assert result['cards'][0]['claim_ids']==[claim['id']]
    assert knowledge['concepts'][0]['claim_ids']==[claim['id']]
    assert knowledge['formulas'][0]['source_ids']==['s1']
    assert knowledge['experiments'][0]['dataset']=='TestSet'
    assert knowledge['visuals'][0]['source_id']=='s1'

def test_legacy_card_response_derives_compatible_knowledge_ir():
    validate=card_validator([{'id':'s1','text':'A fact.'}])
    result=validate({'cards':[{'id':'old','title':'Topic','points':['A fact.'],'source_ids':['s1'],
                               'importance':3,'formulas':[{'latex':'a=b'}],'examples':[]}]})
    assert len(result['knowledge']['claims'])==1
    assert result['knowledge']['claims'][0]['source_ids']==['s1']
    assert result['cards'][0]['claim_ids']==[result['knowledge']['claims'][0]['id']]
    assert result['knowledge']['formulas'][0]['latex']=='a=b'

def test_dangling_model_claim_links_are_rebuilt_from_valid_sources():
    validate=card_validator([{'id':'s1','text':'Result is 7.'}])
    result=validate({'cards':[{'id':'card','title':'Result','points':['Result is 7.'],'source_ids':['s1'],
                               'importance':3,'formulas':[],'examples':[],'claim_ids':['missing']}],
                     'claims':[{'id':'real','statement':'Result is 7.','claim_type':'result','source_ids':['s1'],
                                'conditions':[],'confidence':1,'importance':3}]})
    claim_id=result['knowledge']['claims'][0]['id']
    assert result['cards'][0]['claim_ids']==[claim_id]
    assert result['knowledge']['experiments'][0]['claim_ids']==[claim_id]

def test_cross_batch_normalization_merges_paraphrases_and_repairs_references():
    claims=[
        {'id':'c1','statement':'TestSet accuracy reaches 92%.','claim_type':'result','source_ids':['s1'],
         'conditions':['TestSet'],'confidence':.9,'importance':4},
        {'id':'c2','statement':'Accuracy on TestSet is 92%.','claim_type':'result','source_ids':['s2'],
         'conditions':['TestSet'],'confidence':1,'importance':3},
    ]
    cards=[
        {'id':'a','title':'TestSet evaluation','points':[claims[0]['statement']],'source_ids':['s1'],
         'importance':4,'formulas':[],'examples':[],'claim_ids':['c1']},
        {'id':'b','title':'TestSet evaluation','points':[claims[1]['statement']],'source_ids':['s2'],
         'importance':3,'formulas':[],'examples':[],'claim_ids':['c2']},
    ]
    knowledge={'claims':claims,
        'concepts':[{'id':'x1','name':'Accuracy','definition':'metric','claim_ids':['c1'],'depends_on':[]},
                    {'id':'x2','name':'accuracy','definition':'evaluation metric','claim_ids':['c2'],'depends_on':[]}],
        'formulas':[],
        'experiments':[{'id':'e','name':'Evaluation','dataset':'TestSet','setup':[],'metrics':['accuracy'],
                        'results':['92%'],'source_ids':['s2'],'claim_ids':['c2']}],
        'visuals':[]}
    normalized_cards,normalized=normalize_knowledge(cards,knowledge)
    assert len(normalized['claims'])==1
    assert normalized['claims'][0]['source_ids']==['s1','s2']
    assert len(normalized['concepts'])==1 and len(normalized_cards)==1
    claim_id=normalized['claims'][0]['id']
    assert normalized_cards[0]['claim_ids']==[claim_id]
    assert normalized['experiments'][0]['claim_ids']==[claim_id]

def test_cross_batch_normalization_preserves_different_results_and_conditions():
    base={'claim_type':'result','confidence':1,'importance':4}
    claims=[
        {'id':'c1','statement':'Accuracy is 41.0%.','source_ids':['s1'],'conditions':['dev split'],**base},
        {'id':'c2','statement':'Accuracy is 41.8%.','source_ids':['s2'],'conditions':['dev split'],**base},
        {'id':'c3','statement':'Accuracy is 41.0%.','source_ids':['s3'],'conditions':['test split'],**base},
        {'id':'c4','statement':'Accuracy is not 41.0%.','source_ids':['s4'],'conditions':['dev split'],**base},
    ]
    cards=[{'id':claim['id'],'title':'Result '+claim['id'],'points':[claim['statement']],
            'source_ids':claim['source_ids'],'importance':4,'formulas':[],'examples':[],
            'claim_ids':[claim['id']]} for claim in claims]
    _,normalized=normalize_knowledge(cards,{'claims':claims,'concepts':[],'formulas':[],
                                            'experiments':[],'visuals':[]})
    assert len(normalized['claims'])==4

def test_claim_context_groups_experiment_fields_and_detects_only_same_scope_conflicts():
    context=lambda dataset:{'datasets':[dataset],'models':['Model-v2'],'metrics':['Accuracy'],
                            'setup':['fixed split'],'values':[],'negated':False}
    claims=[
        {'id':'c1','statement':'Accuracy is 91%.','claim_type':'result','source_ids':['s1'],
         'conditions':[],'context':context('TestSet'),'confidence':1,'importance':4},
        {'id':'c2','statement':'Accuracy is 93%.','claim_type':'result','source_ids':['s2'],
         'conditions':[],'context':context('TestSet'),'confidence':1,'importance':4},
        {'id':'c3','statement':'Accuracy is 88%.','claim_type':'result','source_ids':['s3'],
         'conditions':[],'context':context('OtherSet'),'confidence':1,'importance':4},
    ]
    cards=[{'id':c['id'],'title':c['id'],'points':[c['statement']],'source_ids':c['source_ids'],
            'importance':4,'formulas':[],'examples':[],'claim_ids':[c['id']]} for c in claims]
    _,knowledge=normalize_knowledge(cards,{'claims':claims,'concepts':[],'formulas':[],
                                           'experiments':[],'visuals':[]})
    assert len(knowledge['conflicts'])==1
    conflict=knowledge['conflicts'][0]
    assert conflict['kind']=='numeric_mismatch' and set(conflict['claim_ids'])=={'c1','c2'}
    assert 'c3' not in conflict['claim_ids']

def test_card_validation_enriches_claim_context_from_experiment():
    validate=card_validator([{'id':'s1','text':'Model-v2 gets Accuracy 92% on TestSet.'}])
    result=validate({'cards':[{'id':'k','title':'Result','points':['Accuracy 92%'],'source_ids':['s1'],
                               'importance':4,'formulas':[],'examples':[],'claim_ids':['c']}],
                     'claims':[{'id':'c','statement':'Model-v2 gets Accuracy 92% on TestSet.',
                                'claim_type':'result','source_ids':['s1'],'conditions':[],
                                'confidence':1,'importance':4}],
                     'experiments':[{'id':'e','name':'Main','dataset':'TestSet','setup':['fixed split'],
                                     'metrics':['Accuracy'],'results':['92%'],'source_ids':['s1'],
                                     'claim_ids':['c']} ]})
    context=result['knowledge']['claims'][0]['context']
    assert context['datasets']==['TestSet'] and context['metrics']==['Accuracy']
    assert context['setup']==['fixed split'] and context['values']==['92%']

def test_conflict_detection_does_not_mix_model_variants_or_baselines():
    base={'claim_type':'result','conditions':[],'context':{},'confidence':1,'importance':4}
    claims=[
        {'id':'c1','statement':'Transformer (base) BLEU is 27.3.','source_ids':['s'],**base},
        {'id':'c2','statement':'Transformer (big) BLEU is 28.4.','source_ids':['s'],**base},
        {'id':'c3','statement':'HybridRank MRR@10 is 0.386.','source_ids':['s'],**base},
        {'id':'c4','statement':'BM25 MRR@10 is 0.351.','source_ids':['s'],**base},
    ]
    cards=[{'id':c['id'],'title':c['id'],'points':[c['statement']],'source_ids':['s'],'importance':4,
            'formulas':[],'examples':[],'claim_ids':[c['id']]} for c in claims]
    _,knowledge=normalize_knowledge(cards,{'claims':claims,'concepts':[],'formulas':[],
                                           'experiments':[],'visuals':[]})
    assert knowledge['conflicts']==[]

def test_global_concept_graph_uses_all_concepts_and_breaks_prerequisite_cycles():
    class GraphContext:
        def __init__(self): self.payload=None;self.warnings=[]
        def llm(self,role,system,data,validate):
            self.payload=data
            return validate({'edges':[
                {'source_id':'b','target_id':'c','relation':'prerequisite','reason':'B before C'},
                {'source_id':'c','target_id':'a','relation':'prerequisite','reason':'would cycle'},
                {'source_id':'a','target_id':'d','relation':'related','reason':'shared topic'}]})
        def warn(self,message): self.warnings.append(message)
    knowledge={'claims':[],
        'concepts':[{'id':'a','name':'A','definition':'','claim_ids':[],'depends_on':[]},
                    {'id':'b','name':'B','definition':'','claim_ids':[],'depends_on':['a']},
                    {'id':'c','name':'C','definition':'','claim_ids':[],'depends_on':[]},
                    {'id':'d','name':'D','definition':'','claim_ids':[],'depends_on':[]}],
        'formulas':[],'experiments':[],'visuals':[],'conflicts':[]}
    ctx=GraphContext();result=build_concept_graph(ctx,knowledge)
    assert {item['id'] for item in ctx.payload}=={'a','b','c','d'}
    prerequisites={(e['source_id'],e['target_id']) for e in result['concept_edges']
                   if e['relation']=='prerequisite'}
    assert ('a','b') in prerequisites and ('b','c') in prerequisites
    assert ('c','a') not in prerequisites

def test_global_plan_rejects_cross_chapter_duplicate_claim_and_falls_back_together():
    class PlanContext:
        def __init__(self): self.calls=0;self.warnings=[]
        def llm(self,role,system,data,validate):
            self.calls+=1
            return validate({'chapters':[{'id':'x','title':'One','card_ids':['a']},
                                         {'id':'y','title':'Two','card_ids':['b']}]})
        def warn(self,message): self.warnings.append(message)
    cards=[{'id':'a','title':'A','points':['a'],'source_ids':['s1'],'importance':3,
            'formulas':[],'examples':[],'claim_ids':['same']},
           {'id':'b','title':'B','points':['b'],'source_ids':['s2'],'importance':3,
            'formulas':[],'examples':[],'claim_ids':['same']}]
    ctx=PlanContext();chapters=plan(ctx,cards,{'concepts':[],'concept_edges':[]})
    assert ctx.calls==1 and len(chapters)==1
    assert chapters[0]['card_ids']==['a','b'] and ctx.warnings

def test_normalization_assigns_each_claim_to_one_main_card():
    claim={'id':'c','statement':'One fact.','claim_type':'fact','source_ids':['s'],
           'conditions':[],'confidence':1,'importance':4}
    cards=[{'id':'a','title':'First topic','points':['One fact.'],'source_ids':['s'],'importance':4,
            'formulas':[],'examples':[],'claim_ids':['c']},
           {'id':'b','title':'Different application','points':['Apply the fact.'],'source_ids':['s'],
            'importance':3,'formulas':[],'examples':[],'claim_ids':['c']}]
    normalized,_=normalize_knowledge(cards,{'claims':[claim],'concepts':[],'formulas':[],
                                            'experiments':[],'visuals':[]})
    assert sum('c' in card['claim_ids'] for card in normalized)==1
    assert next(card for card in normalized if card['claim_ids'])['title']=='First topic'

