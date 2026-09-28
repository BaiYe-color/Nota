"""Native document extraction, selective vision and per-page checkpoints."""
import base64
import json
import os
from collections import Counter
from pathlib import Path
import pymupdf
from contracts import SourceBlock
from engine import Cancelled, digest, atomic_json, MODELS
from asr import transcribe_local_file, semantic_segments, PARAFORMER_MODEL, BASE

AUDIO={'.mp3','.wav','.m4a','.flac','.aac','.ogg','.opus','.wma'}
IMAGES={'.png','.jpg','.jpeg','.bmp','.webp','.tiff','.tif'}
DOCS={'.pdf','.docx','.doc','.pptx','.ppt','.txt','.md'}
ACCEPTED=AUDIO|IMAGES|DOCS
VISION_PROMPT='''识别学习材料页面，忠实保留文字、数字、条件、公式和表格。图表需解释其实际内容；不能看清的地方写[待核对]，不要猜测。已有原生文本供校对。返回 {"text":"完整内容，数学使用LaTeX，表格使用Markdown","uncertain":false}。'''
FIGURE_PROMPT='''这是从学习材料中精确裁剪出的图、表或结构图，不是整页截图。结合附近原文说明它表达什么，以及读者应从图中读出什么。保留图号、表号、变量、数字和图例；不要复述无关页面正文，不要猜测看不清的内容。第一行写简短标题，随后用2至6条要点解释。返回 {"text":"Markdown说明","uncertain":false}。'''

def _expanded(rect,page_rect,pad=8):
    return pymupdf.Rect(max(page_rect.x0,rect.x0-pad),max(page_rect.y0,rect.y0-pad),
                        min(page_rect.x1,rect.x1+pad),min(page_rect.y1,rect.y1+pad))

def pdf_visual_regions(page,image_frequency,total):
    """Return meaningful figure/table regions; never use a born-digital full page as an asset."""
    candidates=[]
    page_area=page.rect.get_area()
    image_rects=[]
    for info in page.get_image_info(xrefs=True):
        rect=pymupdf.Rect(info.get('bbox') or ())
        if rect.is_empty or rect.get_area()<page_area*.008:
            continue
        if image_frequency.get(info.get('xref'),1)>=max(3,total*.6):
            continue
        image_rects.append(_expanded(rect,page.rect,10))
    if len(image_rects)>1:
        merged=pymupdf.Rect(image_rects[0])
        for rect in image_rects[1:]: merged|=rect
        same_band=max(r.y0 for r in image_rects)<min(r.y1 for r in image_rects)
        if same_band and merged.get_area()<page_area*.5:
            # Side-by-side subfigures often keep titles above and a shared caption below
            # the embedded image boxes. Preserve those labels without falling back to a page shot.
            merged=pymupdf.Rect(max(page.rect.x0,merged.x0-70),max(page.rect.y0,merged.y0-22),
                                min(page.rect.x1,merged.x1+40),min(page.rect.y1,merged.y1+30))
            candidates.append(('figure',merged))
        else:
            candidates.extend(('figure',rect) for rect in image_rects)
    else:
        candidates.extend(('figure',rect) for rect in image_rects)
    try:
        for table in page.find_tables().tables:
            rect=_expanded(pymupdf.Rect(table.bbox),page.rect,8)
            if rect.get_area()>=page_area*.025:
                candidates.append(('table',rect))
    except Exception:
        pass
    drawings=page.get_drawings()
    if len(drawings)>50:
        rect=pymupdf.Rect(drawings[0]['rect'])
        for drawing in drawings[1:]: rect|=drawing['rect']
        rect=_expanded(rect,page.rect,10)
        if rect.get_area()>=page_area*.025:
            candidates.append(('figure',rect))
    # Prefer the larger region when table detection and vector drawings describe the same graphic.
    result=[]
    for kind,rect in sorted(candidates,key=lambda item:item[1].get_area(),reverse=True):
        overlap=False
        for _,existing in result:
            intersection=rect & existing
            if intersection.get_area()>=min(rect.get_area(),existing.get_area())*.72:
                overlap=True; break
        if not overlap: result.append((kind,rect))
    return result[:4]

LAYOUT_PROMPT = """识别这一页视觉学习材料的内容与版面。它可能是课堂课件、扫描讲义、教材页或个人笔记。只返回区域数组，不重复输出整页全文：
{"uncertain":false,"regions":[{"role":"title","bbox":[0.1,0.05,0.9,0.2],"text":"标题"},{"role":"text","bbox":[0.1,0.3,0.9,0.8],"text":"正文"}]}。
role只能是title、text、table、formula、diagram、code、footer。bbox必须为0到1归一化的[左,上,右,下]，每区恰好四个数。每个区域必须有text字符串。按阅读顺序保留完整正文、数字、条件、否定和提问。表格只在table区域用Markdown表格转录一次，不能将每个单元格拆成区域。公式使用实际LaTeX。diagram指独立图形、流程或多要素关系结构，text保留节点、连线和关键数据；纯文字框不能标diagram。页眉页脚放footer；整页不能标diagram；不重复重叠内容。不要讲解，不推导，不补全看不清的文字；看不清时标uncertain。一般3至10个区域即可。"""

# Kept as an import-compatible name for callers/tests from the first layout rollout.
SLIDE_PROMPT = LAYOUT_PROMPT

_SLIDE_ROLES={
    'heading':'title','header':'title','subtitle':'title','body':'text','paragraph':'text',
    'content':'text','image':'diagram','figure':'diagram','chart':'diagram','flowchart':'diagram',
    'math':'formula','equation':'formula','bottom':'footer','footnote':'footer',
}

def _slide_box(raw):
    """Normalize the common region formats emitted by different vision providers."""
    if isinstance(raw,dict):
        left=raw.get('left',raw.get('x',raw.get('x0')))
        top=raw.get('top',raw.get('y',raw.get('y0')))
        right=raw.get('right',raw.get('x1'))
        bottom=raw.get('bottom',raw.get('y1'))
        width,height=raw.get('width'),raw.get('height')
        if right is None and left is not None and width is not None: right=left+width
        if bottom is None and top is not None and height is not None: bottom=top+height
        raw=[left,top,right,bottom]
    if not isinstance(raw,(list,tuple)) or len(raw)!=4 or not all(isinstance(x,(int,float)) for x in raw):
        return None
    box=[float(x) for x in raw]
    # Qwen's layout responses commonly use a 0..1000 grid.  Percentages are
    # also unambiguous here.  Do not guess arbitrary pixel dimensions.
    extent=max(abs(x) for x in box)
    if extent>1 and extent<=100: box=[x/100 for x in box]
    elif extent>100 and extent<=1000: box=[x/1000 for x in box]
    return box

def validate_slide(value):
    if not isinstance(value,dict): raise ValueError('课件返回的不是 JSON 对象')
    regions=value.get('regions',value.get('layout',value.get('blocks',value.get('elements',[]))))
    if not isinstance(regions,list) or not regions:raise ValueError('课件必须返回页面区域')
    normalized=[]
    for region in regions:
        if not isinstance(region,dict):
            raise ValueError('页面区域必须是对象')
        role=str(region.get('role',region.get('type',region.get('kind','')))).lower()
        role=_SLIDE_ROLES.get(role,role)
        box=_slide_box(region.get('bbox',region.get('coordinates',region.get('box',region.get('position',[])))))
        if role not in {'title','text','table','formula','diagram','code','footer'} or not box:
            raise ValueError('页面区域类型或坐标无效')
        if not all(isinstance(x,(int,float)) and 0<=x<=1 for x in box) or box[0]>=box[2] or box[1]>=box[3]:
            raise ValueError('页面区域坐标必须归一化')
        text=region.get('text',region.get('content',region.get('label',region.get('description'))))
        if not isinstance(text,str):
            raise ValueError('每个regions元素必须有text字符串，包括diagram应填写节点和连线描述。示例：{"role":"diagram","bbox":[0.1,0.2,0.8,0.7],"text":"文档→分词→排序"}')
        normalized.append({'role':role,'bbox':[round(x,6) for x in box],'text':text})
    text='\n\n'.join(r['text'] for r in normalized if r['role']!='footer').strip()
    if not text:raise ValueError('课件识别为空')
    return {'text':text,'uncertain':bool(value.get('uncertain',False)),'regions':normalized}

def extract(ctx,files,options):
    blocks=[]; counts=[]
    for index,file in enumerate(files):
        ctx.check()
        path=Path(file['path']); ext=path.suffix.lower(); doc=file['id']
        pages=[]
        def progress(message): ctx.progress(.04+.40*index/len(files),file['name']+' · '+message)
        def vision(native,png):
            result=ctx.llm('vision',VISION_PROMPT,{'native_text':native},image='data:image/png;base64,'+base64.b64encode(png).decode(),
                           validate=lambda x: {'text':str(x['text']),'uncertain':bool(x.get('uncertain',False))})
            if not result['text'].strip(): raise ValueError('视觉识别返回空白')
            return result
        def figure_vision(native,png):
            result=ctx.llm('vision',FIGURE_PROMPT,{'nearby_text':native},image='data:image/png;base64,'+base64.b64encode(png).decode(),
                           validate=lambda x: {'text':str(x['text']),'uncertain':bool(x.get('uncertain',False))})
            if not result['text'].strip(): raise ValueError('图表识别返回空白')
            return result
        def asset(png):
            key=digest([file['sha'],base64.b64encode(png).decode()])+'.png'
            target=ctx.store.root/'artifacts'/key
            if not target.exists(): target.write_bytes(png)
            return key
        def one(number,native,png_fn=None,visual=False):
            ctx.progress(.04+.40*(index+(number-1)/max(1,limit))/len(files),file['name']+f' · 解析第 {number}/{limit} 页/单元')
            params={'sha':file['sha'],'page':number,'handwritten':options.handwritten,'mode':options.visual_mode,
                    'images':options.include_images,'model':MODELS['vision'],'provider':os.getenv('OPENAI_BASE_URL','https://api.openai.com/v1'),'fallback':os.getenv('NOTA_VISION_FALLBACK_MODEL','qwen-vl-plus'),'fallback_provider':os.getenv('QWEN_BASE_URL','https://dashscope.aliyuncs.com/compatible-mode/v1'),'parser':5}
            def run():
                needs=png_fn and (options.handwritten or options.visual_mode=='always' or len(native.strip())<40 or visual)
                png=png_fn() if needs or (options.include_images and visual and png_fn) else None
                result=vision(native,png) if needs else {'text':native,'uncertain':False}
                # A handwritten page is evidence for OCR, not a publishable
                # illustration.  Exposing the raster as `asset` made every
                # note contain an unrelated full-page scan.  Keep it only for
                # the source panel; later visual extraction may add verified
                # local crops as separate assets.
                page_asset=asset(png) if png and options.include_images and (visual or len(native.strip())<40) else None
                return {**result,'page':number,'method':'vision' if needs else 'native',
                        'asset':None if options.handwritten else page_asset,
                        'evidence_asset':page_asset if options.handwritten else None}
            try: pages.append(ctx.cached('parse-page',params,run))
            except Cancelled: raise
            except Exception as exc: ctx.warn(f"{file['name']} 第 {number} 页失败：{type(exc).__name__}；可重试")
        if ext in AUDIO:
            asr_route=ctx.model_config.get('asr',{})
            params={'sha':file['sha'],'diarization':options.diarization,'model':asr_route.get('model') or PARAFORMER_MODEL,'base':asr_route.get('base_url') or BASE}
            task_path=ctx.store.root/'cache'/'asr-tasks'/(digest(params)+'.json')
            def transcribe():
                task=json.loads(task_path.read_text())['task_id'] if task_path.exists() else None
                try:
                    return transcribe_local_file(path,enable_diarization=options.diarization,progress_cb=progress,check_cancel=ctx.check,
                        task_id=task,on_task=lambda tid:atomic_json(task_path,{'task_id':tid}),config=asr_route)
                except RuntimeError:
                    task_path.unlink(missing_ok=True)
                    raise
            sentences=ctx.cached('asr',params,transcribe)
            segments=semantic_segments(sentences)
            for n,seg in enumerate(segments):
                speaker=f"说话人{seg['speaker_id']}: " if options.diarization and seg.get('speaker_id') is not None else ''
                blocks.append(SourceBlock(id=f'{doc[:12]}-t{n+1}',document_id=doc,document_name=file['name'],
                    start_ms=seg['start_ms'],end_ms=seg['end_ms'],text=speaker+seg['text'],kind='audio',method='asr',
                    speaker_id=seg.get('speaker_id'),asr_sentences=seg['sentences'],
                    material_type=file.get('material_type','general'),material_role=file.get('material_role','primary')).model_dump())
            counts.append({'document':file['name'],'total':len(segments),'processed':len(segments),'unit':'语音片段'})
            continue
        if ext in ('.doc','.ppt'):
            from normalize import normalize_to_pdf
            converted=ctx.store.root/'cache'/'converted'/file['sha']
            converted.mkdir(parents=True,exist_ok=True)
            path=Path(normalize_to_pdf(str(path),lambda _:converted))
            ext='.pdf'
        if ext=='.pdf' or ext in IMAGES:
            with pymupdf.open(path) as original:
                if ext in IMAGES:
                    document=pymupdf.open('pdf',original.convert_to_pdf())
                else: document=original
                try:
                    total=len(document); limit=min(total,options.max_pages or total)
                    image_frequency=Counter(im[0] for p in document for im in p.get_images())
                    for n in range(limit):
                        page=document[n]
                        native=page.get_text('text',sort=True)
                        # Scanned pages need full-page vision. Born-digital PDFs keep native text and
                        # create separate cropped visual evidence for figures and tables.
                        scanned=len(native.strip())<40
                        # Layout extraction is about the page's visual form, not the user's
                        # writing preference.  A scanned teacher slide may be labelled
                        # "personal notes" for prose style; it must still never turn into a
                        # full-page published image.  Handwritten mode is intentionally
                        # excluded because its free-form annotations need the handwriting OCR
                        # prompt rather than rigid layout regions.
                        if scanned and not options.handwritten:
                            png=page.get_pixmap(dpi=150,alpha=False).tobytes('png')
                            params={'sha':file['sha'],'page':n+1,'layout_schema':3,'model':MODELS['vision']}
                            data=ctx.cached('slide-layout',params,lambda p=png,t=native:ctx.llm(
                                'vision',LAYOUT_PROMPT,{'native_text':t},image='data:image/png;base64,'+base64.b64encode(p).decode(),validate=validate_slide))
                            data=validate_slide(data)
                            evidence=asset(png)
                            pages.append({'text':data['text'],'uncertain':data['uncertain'],'page':n+1,
                                          'method':'slide-layout','asset':None,'evidence_asset':evidence,
                                          'page_layout':data['regions'],'kind':'text'})
                            if options.include_images:
                                for region in data['regions']:
                                    x0,y0,x1,y1=region['bbox']
                                    if region['role'] not in {'diagram','table','formula','code'} or (x1-x0)*(y1-y0)>.75:continue
                                    rect=pymupdf.Rect(x0*page.rect.width,y0*page.rect.height,x1*page.rect.width,y1*page.rect.height)
                                    crop=page.get_pixmap(dpi=180,clip=rect,alpha=False).tobytes('png')
                                    pages.append({'text':region['text'],'page':n+1,'method':'layout-region','kind':'visual',
                                                  'asset':asset(crop),'evidence_asset':evidence,'bbox':region['bbox'],
                                                  'layout_role':region['role'],'uncertain':data['uncertain']})
                            ctx.progress(.04+.40*(index+(n+1)/limit)/len(files),file['name']+f' · 结构化解析 {n+1}/{limit}')
                            continue
                        one(n+1,native,lambda p=page:p.get_pixmap(dpi=160).tobytes('png'),False)
                        if ext=='.pdf' and not scanned:
                            for region_no,(kind,rect) in enumerate(pdf_visual_regions(page,image_frequency,total),1):
                                ctx.check()
                                png=page.get_pixmap(dpi=180,clip=rect,alpha=False).tobytes('png')
                                nearby=page.get_textbox(_expanded(rect,page.rect,70)).strip()
                                params={'sha':file['sha'],'page':n+1,'region':region_no,
                                        'bbox':[round(v,2) for v in rect],'model':MODELS['vision'],'parser':5}
                                try:
                                    data=ctx.cached('pdf-visual-region',params,lambda p=png,t=nearby:figure_vision(t,p))
                                    pages.append({**data,'page':n+1,'method':'vision-region','kind':'table' if kind=='table' else 'visual',
                                                  'asset':asset(png) if options.include_images else None})
                                except Cancelled: raise
                                except Exception as exc:
                                    ctx.warn(f"{file['name']} 第 {n+1} 页图表 {region_no} 识别失败：{type(exc).__name__}")
                finally:
                    if document is not original: document.close()
        elif ext=='.pptx':
            from pptx import Presentation
            prs=Presentation(path); total=len(prs.slides); limit=min(total,options.max_pages or total)
            for n,slide in enumerate(list(prs.slides)[:limit]):
                ctx.check(); text=[]; pictures=[]; layout=[]
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text.append(shape.text)
                        layout.append({'role':'title' if shape==slide.shapes.title else 'text','text':shape.text,
                                       'bbox':[shape.left/prs.slide_width,shape.top/prs.slide_height,
                                               (shape.left+shape.width)/prs.slide_width,(shape.top+shape.height)/prs.slide_height]})
                    if shape.has_table:
                        text.append('\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows))
                    if hasattr(shape,'image'):
                        try:
                            image_doc=pymupdf.open(stream=shape.image.blob)
                            pixels=image_doc[0].get_pixmap().tobytes('png'); image_doc.close()
                            pictures.append(pixels)
                        except Exception: ctx.warn(f"{file['name']} 第 {n+1} 页有无法读取的图片")
                    if getattr(shape,'has_chart',False):
                        chart=shape.chart
                        try:
                            text.append('图表：'+str([c.label for c in chart.plots[0].categories]))
                            text.extend(str(s.name)+': '+str(list(s.values)) for s in chart.series)
                        except Exception: ctx.warn(f"{file['name']} 第 {n+1} 页图表需人工核对")
                if slide.has_notes_slide:
                    text.append('讲者备注：'+slide.notes_slide.notes_text_frame.text)
                one(n+1,'\n'.join(text))
                if pages and pages[-1]['page']==n+1:pages[-1]['page_layout']=layout
                for v,png in enumerate(pictures):
                    try:
                        d=ctx.cached('ppt-visual',{'sha':file['sha'],'page':n,'image':v,'model':MODELS['vision']},lambda p=png:vision('',p))
                        pages.append({**d,'page':n+1,'method':'vision','asset':asset(png) if options.include_images else None})
                    except Cancelled: raise
                    except Exception: ctx.warn(f"{file['name']} 第 {n+1} 页图片 {v+1} 识别失败")
            ctx.warn(f"{file['name']} 使用原生 PPT 提取；复杂组合图形请导出 PDF 后使用视觉模式核对")
        elif ext=='.docx':
            from docx import Document
            docx=Document(path); text=[]
            for item in docx.iter_inner_content():
                if hasattr(item,'text'): text.append(item.text)
                else: text.append('\n'.join(' | '.join(c.text for c in row.cells) for row in item.rows))
            total=1; limit=1
            one(1,'\n'.join(text))
            for v,rel in enumerate(docx.part.rels.values()):
                if not rel.reltype.endswith('/image') or rel.is_external: continue
                try:
                    with pymupdf.open(stream=rel.target_part.blob) as image_doc:
                        png=image_doc[0].get_pixmap().tobytes('png')
                    d=ctx.cached('docx-visual',{'sha':file['sha'],'image':v,'model':MODELS['vision']},lambda p=png:vision('',p))
                    pages.append({**d,'page':1,'method':'vision','asset':asset(png) if options.include_images else None})
                except Cancelled: raise
                except Exception: ctx.warn(f"{file['name']} 嵌入图片 {v+1} 识别失败，可重试")
        else:
            total=1; limit=1
            one(1,path.read_text(encoding='utf-8-sig'))
        processed=len({p['page'] for p in pages})
        counts.append({'document':file['name'],'total':total,'processed':processed,'unit':'页/原生文档单元'})
        if limit<total: ctx.warn(f"{file['name']} 仅处理前 {limit}/{total} 页")
        for n,p in enumerate(pages):
            value=p['text'].strip()
            if not value:
                ctx.warn(f"{file['name']} 第 {p['page']} 页没有提取到文字")
                continue
            # Split large native sections to keep every model input bounded.
            for part,start in enumerate(range(0,len(value),3500)):
                blocks.append(SourceBlock(id=f'{doc[:12]}-p{p["page"]}-b{n}-{part}',document_id=doc,document_name=file['name'],
                    page=p['page'],text=value[start:start+3500],kind=p.get('kind','visual' if p.get('asset') else 'text'),
                    method=p['method'],asset=p.get('asset') if part==0 else None,uncertain=p.get('uncertain',False),
                    evidence_asset=p.get('evidence_asset'),bbox=p.get('bbox'),page_layout=p.get('page_layout',[]),
                    material_type=file.get('material_type','general'),material_role=file.get('material_role','primary')).model_dump())
    if not blocks: raise RuntimeError('没有成功提取到内容；请检查失败页后重试')
    return blocks,counts
