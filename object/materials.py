"""Material type detection, role resolution, and isolated style sampling."""
from __future__ import annotations

import re
from pathlib import Path

import pymupdf

from contracts import MaterialInput, Options


AUDIO={'.mp3','.wav','.m4a','.flac','.aac','.ogg','.opus','.wma'}


def detect_material_type(file: dict, options: Options) -> tuple[str,float,str]:
    path=Path(file['path']); ext=path.suffix.lower(); name=file['name'].lower()
    patterns=[
        ('personal_notes',r'笔记|notes?',.9),('paper',r'论文|paper|study|arxiv',.86),
        ('textbook',r'教材|讲义|chapter|chap\d',.82),('exercise',r'习题|试卷|exercise|exam',.9),
        ('meeting',r'会议|纪要|meeting',.9)]
    for kind,pattern,confidence in patterns:
        if re.search(pattern,name,re.I): return kind,confidence,'根据文件名识别'
    if options.handwritten and ext in {'.pdf','.png','.jpg','.jpeg','.bmp','.webp','.tiff','.tif'}:
        return 'personal_notes',.85,'启用了手写材料选项'
    if ext in {'.ppt','.pptx'}: return 'lecture_slides',.99,'PowerPoint 文件'
    if ext in AUDIO: return 'meeting',.68,'录音默认按会议/课堂发言处理'
    if ext=='.pdf':
        try:
            with pymupdf.open(path) as doc:
                text='\n'.join(page.get_text('text',sort=True) for page in list(doc)[:3])
                pages=len(doc)
            if re.search(r'\babstract\b|\breferences\b|\bdoi\b|摘要|参考文献',text,re.I):
                return 'paper',.88,'检测到论文结构词'
            if pages>=5 and len(text.strip())<2400:
                return 'lecture_slides',.62,'页面较多且单页文字较少'
            if re.search(r'第.{0,4}章|chapter|练习题|例题',text,re.I):
                return 'textbook',.72,'检测到教材章节结构'
        except Exception:
            pass
        return 'general',.45,'PDF 类型特征不足'
    if ext=='.docx': return 'general',.58,'Word 文档类型特征不足'
    return 'general',.5,'未检测到明确类型特征'


def resolve_materials(files: list[dict], inputs: list[dict], options: Options) -> list[dict]:
    requested={item['file_id']:MaterialInput.model_validate(item) for item in inputs}
    resolved=[]
    for file in files:
        item=requested.get(file['id'],MaterialInput(file_id=file['id']))
        detected,confidence,reason=detect_material_type(file,options)
        effective=detected if item.source_type=='auto' else item.source_type
        if item.role=='style_reference' and Path(file['path']).suffix.lower() in AUDIO:
            raise ValueError('录音暂不能作为风格参考；请将其设为主要内容或补充资料')
        resolved.append({
            'file_id':file['id'],'name':file['name'],'requested_type':item.source_type,
            'detected_type':detected,'detection_confidence':confidence,'detection_reason':reason,
            'effective_type':effective,'role':item.role})
    return resolved


def style_sample(file: dict, limit=5000) -> str:
    path=Path(file['path']); ext=path.suffix.lower(); parts=[]
    try:
        if ext in {'.txt','.md'}:
            return path.read_text(encoding='utf-8-sig')[:limit]
        if ext=='.pdf':
            with pymupdf.open(path) as doc:
                return '\n'.join(page.get_text('text',sort=True) for page in list(doc)[:5])[:limit]
        if ext=='.docx':
            from docx import Document
            doc=Document(path)
            for item in doc.iter_inner_content():
                if hasattr(item,'text'): parts.append(item.text)
                else: parts.extend(' | '.join(cell.text for cell in row.cells) for row in item.rows)
        elif ext=='.pptx':
            from pptx import Presentation
            for slide in Presentation(path).slides:
                parts.extend(shape.text for shape in slide.shapes if getattr(shape,'has_text_frame',False))
        return '\n'.join(parts)[:limit]
    except Exception:
        return ''


def apply_materials(files: list[dict], resolved: list[dict]) -> list[dict]:
    by_id={item['file_id']:item for item in resolved}
    return [{**file,'material_type':by_id[file['id']]['effective_type'],'material_role':by_id[file['id']]['role']}
            for file in files]
