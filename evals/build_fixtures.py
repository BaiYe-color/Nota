"""Build deterministic, non-private files used by the Nota quality suite."""
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(__file__).resolve().parent
FIXTURES = ROOT / "fixtures"
CHINESE_FONT = Path(r"C:\Windows\Fonts\msyh.ttc")


def add_text(page, rect, text, size=15, latin=False):
    if latin:
        page.insert_textbox(rect, text, fontsize=size, fontname="helv", lineheight=1.35)
    elif CHINESE_FONT.exists():
        page.insert_textbox(rect, text, fontsize=size, fontname="msyh", fontfile=str(CHINESE_FONT), lineheight=1.35)
    else:
        page.insert_textbox(rect, text, fontsize=size, fontname="china-s", lineheight=1.35)


def build_paper():
    path = FIXTURES / "hybrid-retrieval-study.pdf"
    path.unlink(missing_ok=True)
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    add_text(page, pymupdf.Rect(45, 45, 550, 100), "HybridRank: A Controlled Study of Hybrid Retrieval", 20, True)
    add_text(page, pymupdf.Rect(45, 115, 550, 430),
             "Abstract\nHybridRank combines BM25 sparse retrieval with dense vector retrieval. "
             "The evaluation uses 1,200 Chinese support queries. HybridRank reaches MRR@10 0.412, "
             "compared with 0.386 for the dense-only baseline. Median retrieval latency is 42 ms.\n\n"
             "Method\nScores are normalized separately and combined with weight 0.35 for BM25 and "
             "0.65 for dense retrieval. The reranker is not used in this controlled experiment.", latin=True)
    page = doc.new_page(width=595, height=842)
    add_text(page, pymupdf.Rect(45, 50, 550, 360),
             "Results\nMethod | MRR@10 | Median latency\nDense only | 0.386 | 31 ms\n"
             "BM25 only | 0.351 | 18 ms\nHybridRank | 0.412 | 42 ms\n\n"
             "Limitations\nThe experiment uses one Chinese support dataset and does not establish "
             "performance on multilingual or open-domain retrieval.", latin=True)
    doc.save(path)
    doc.close()


def image_only_pdf(path, pages):
    path.unlink(missing_ok=True)
    native = pymupdf.open()
    for title, body in pages:
        page = native.new_page(width=595, height=842)
        add_text(page, pymupdf.Rect(45, 55, 550, 105), title, 21)
        add_text(page, pymupdf.Rect(45, 125, 550, 760), body, 16)
    scanned = pymupdf.open()
    for page in native:
        pix = page.get_pixmap(dpi=170, alpha=False)
        target = scanned.new_page(width=page.rect.width, height=page.rect.height)
        target.insert_image(target.rect, stream=pix.tobytes("png"))
    scanned.save(path)
    native.close()
    scanned.close()


def build_scans():
    image_only_pdf(FIXTURES / "binary-search-scan.pdf", [
        ("二分查找实验", "前提：输入数组必须按升序排列。示例数组为 [2, 4, 8, 16, 32]，目标值为 16。"),
        ("算法结论", "使用闭区间 low = 0、high = n - 1。目标 16 的零基索引为 3。"
                     "目标不存在时返回 -1，时间复杂度为 O(log n)。")])
    image_only_pdf(FIXTURES / "database-key-scan.pdf", [
        ("关系模式与候选键", "关系 Enrollment(student_id, course_id, grade)。"
                         "候选键由 student_id 与 course_id 共同组成，二者缺一不可。"),
        ("第二范式检查", "grade 依赖完整组合键 (student_id, course_id)，不存在对组合键一部分的部分依赖。"
                       "因此该关系满足第二范式。")])


def build_slides():
    path = FIXTURES / "project-planning-slides.pptx"
    path.unlink(missing_ok=True)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "笔记项目内部试点"
    slide.placeholders[1].text = "已决定：10 月 8 日启动内部试点\n预算：3,000 元\n暂不公开发布"
    slide.notes_slide.notes_text_frame.text = "手机应用仍是建议，没有立项。"
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "验收指标"
    table = slide.shapes.add_table(4, 2, Inches(1), Inches(2), Inches(7), Inches(2.2)).table
    for r, values in enumerate([
        ["指标", "目标"], ["事实召回率", "至少 95%"], ["审校误报率", "低于 5%"], ["负责人", "小陈"]]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    prs.save(path)


def build_docx():
    path = FIXTURES / "training-handout.docx"
    path.unlink(missing_ok=True)
    doc = Document()
    doc.add_heading("分类模型训练讲义", 0)
    doc.add_paragraph("数据集按 8:2 划分为训练集和测试集。测试集不得参与参数训练。")
    doc.add_heading("训练配置", 1)
    table = doc.add_table(rows=4, cols=2)
    for r, values in enumerate([
        ["配置", "取值"], ["优化器", "Adam"], ["学习率", "0.001"], ["训练轮数", "50"]]):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    doc.add_paragraph("本实验报告准确率与宏平均 F1；不得只报告训练集准确率。")
    doc.save(path)


def main():
    FIXTURES.mkdir(parents=True, exist_ok=True)
    build_paper()
    build_scans()
    build_slides()
    build_docx()
    print("Evaluation fixtures ready:", FIXTURES)


if __name__ == "__main__":
    main()
